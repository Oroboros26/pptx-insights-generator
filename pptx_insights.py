"""
pptx_insights.py — entry point and main orchestrator.

This file is intentionally slim. All logic lives in the modules/ folder.
Run this file directly — never run the module files directly.

Usage:
  python pptx_insights.py --input input/FILE.pptx --output output/FILE.pptx

Pipeline (in order):
  1. load_context()               concatenate context/*.md into the system prompt
  2. Per slide: extract image → GPT-vision analysis → L2 rule check → L3 AI judge
  3. write_notes()                embed insight text into PPTX speaker notes
  4. generate_executive_summary() one synthesis call across all slide insights
  5. generate_pdf_report()        A4 PDF report
  6. send_report_email()          PDF + PPTX attached, via SendGrid or SMTP
"""

import argparse          # CLI argument parsing (--input, --output)
import os                # read env vars and build file paths
import random            # random sampling for the 20% Level 3 spot-check
from pathlib import Path          # resolve .env path relative to this script
from dotenv import load_dotenv    # load secrets from .env into os.environ
from pptx import Presentation    # read and write the PPTX file
from openai import AzureOpenAI   # Azure-hosted GPT client (vision + text)

load_dotenv(Path(__file__).parent / ".env")

from modules.config import SAMPLE_RATE
from modules.context import load_context, build_system_prompt
from modules.slides import extract_slide_image, analyze_slide, write_notes
from modules.evaluation import validate_insight, judge_insight, print_quality_report
from modules.summary import generate_executive_summary
from modules.pdf_report import generate_pdf_report
from modules.email_report import send_report_email


def process_pptx(input_path, output_path):
    """Main pipeline: load PPTX → analyse every slide → save → PDF → email."""

    # One shared Azure OpenAI client reused for all API calls in this run
    client = AzureOpenAI(
        api_key=os.environ.get("AZURE_OPENAI_KEY"),
        azure_endpoint=os.environ["AZURE_OPENAI_ENDPOINT"],  # e.g. https://<your-resource>.openai.azure.com/
        api_version="2025-04-01-preview"
    )

    context       = load_context()
    system_prompt = build_system_prompt(context)
    print(f"Context loaded: {len(context)} chars\n")

    prs = Presentation(input_path)
    print(f"Found {len(prs.slides)} slides\n")

    slide_results = []

    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        print(f"Processing slide {slide_num}...")

        try:
            b64 = extract_slide_image(slide)
            if b64:
                insight = analyze_slide(client, b64, slide_num, system_prompt)
            else:
                insight = f"Slide {slide_num}\nNo image found."

            # Level 2 — fast rule-based quality gate
            l2_flags  = validate_insight(insight)
            l2_passed = len(l2_flags) == 0

            if l2_passed:
                print(f"    L2 ✅ passed")
            else:
                for flag in l2_flags:
                    print(f"    L2 ⚠️  {flag}")

            # Level 3 — AI judge: always on L2 failure, randomly on SAMPLE_RATE % of passes
            l3_result = None
            if b64 and (not l2_passed or random.random() < SAMPLE_RATE):
                reason = "L2 failure" if not l2_passed else f"random {int(SAMPLE_RATE * 100)}% sample"
                print(f"    L3 🧑‍⚖️  judging ({reason})...")
                l3_result = judge_insight(client, b64, insight, slide_num)
                print(f"    L3 scores → accuracy: {l3_result['accuracy']}/5 | format: {l3_result['format']}/5 | tone: {l3_result['tone']}/5")
                if l3_result.get("flags"):
                    for flag in l3_result["flags"]:
                        print(f"    L3 ⚠️  {flag}")

        except Exception as e:
            print(f"    ❌ error on slide {slide_num}: {e} — skipping")
            insight   = f"Slide {slide_num}\nProcessing error — see run logs."
            l2_flags  = [f"processing error: {e}"]
            l2_passed = False
            l3_result = None

        write_notes(slide, insight)
        slide_results.append({
            "slide_num": slide_num,
            "insight":   insight,
            "l2_flags":  l2_flags,
            "l2_passed": l2_passed,
            "l3_result": l3_result,
        })
        print(f"    ✅ done")

    prs.save(output_path)
    print(f"\n[PPTX] Saved: {output_path}")
    print_quality_report(slide_results, SAMPLE_RATE)

    print("\n[Summary] Generating executive summary...")
    executive_summary = generate_executive_summary(client, slide_results)

    pdf_path = output_path.rsplit(".pptx", 1)[0] + "_insights.pdf"
    generate_pdf_report(slide_results, os.path.basename(output_path), pdf_path, executive_summary)
    send_report_email(pdf_path, output_path, executive_summary, slide_results, os.path.basename(output_path))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Add AI-written insight notes to a dashboard PPTX export")
    parser.add_argument("--input",  required=True, help="Path to input PPTX  (e.g. input/file.pptx)")
    parser.add_argument("--output", required=True, help="Path to output PPTX (e.g. output/file.pptx)")
    args = parser.parse_args()
    process_pptx(args.input, args.output)
